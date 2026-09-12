"""Build the raw scientific handoff for the held-out ISIC 2019 test cohort.

The primary population is intentionally restricted to ``ISIC2019/test``.  It
is absent from the original train and validation rows used to fit the 1a heads.
The original validation split is not reused as evaluation data because it was
involved in model selection and threshold calibration.

Run from the repository root::

    python -m derm.analysis.build_isic2019_eval_workbook
"""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path
from typing import Any, Iterable

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from scipy.stats import norm
from sklearn.metrics import (
    accuracy_score,
    average_precision_score,
    balanced_accuracy_score,
    brier_score_loss,
    confusion_matrix,
    f1_score,
    log_loss,
    precision_score,
    recall_score,
    roc_auc_score,
)

from derm import paths


PREDICTIONS_INPUT = paths.INTERNAL_PREDICTIONS
REGISTRY_INPUT = paths.INTERNAL_EXPERIMENT_REGISTRY
METADATA_INPUT = paths.INTERNAL_METADATA
RAW_MANIFEST_INPUT = paths.INTERNAL_RAW_MANIFEST
MODEL_MANIFEST_INPUT = paths.APP_MODELS / "model_manifest.json"
PERFORMANCE_REGISTRY_INPUT = paths.PAPER_PERFORMANCE_REGISTRY
CURRENT_FEATURE_IMPORTANCE_INPUT = paths.CURRENT_ONNX_IMPORTANCE
LEGACY_FEATURE_IMPORTANCE_INPUT = paths.FEATURE_IMPORTANCE_CSV
HANDCRAFT_ABLATION_INPUT = paths.ABLATION_HANDCRAFT_FACTORIAL

PREDICTIONS_OUTPUT = paths.ISIC2019_PREDICTIONS
METRICS_OUTPUT = paths.ISIC2019_METRICS
AUDIT_OUTPUT = paths.ISIC2019_AUDIT
WORKBOOK_OUTPUT = paths.ISIC2019_WORKBOOK

CLASS_ORDER = ("nevus", "melanoma", "atypical")
CLASS_IDS = {name: index for index, name in enumerate(CLASS_ORDER)}
PROBABILITY_COLUMNS = [f"p_{name}" for name in CLASS_ORDER]
BOOTSTRAP_REPLICATES = 2_000
BOOTSTRAP_SEED = 42

NAVY = "17365D"
BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
TEAL = "0F6B78"
LIGHT_TEAL = "DDEBF7"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
ORANGE = "FCE4D6"
RED = "F4CCCC"
WHITE = "FFFFFF"
THIN_GREY = Side(style="thin", color="D9E1F2")


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def relative(path: Path) -> str:
    return path.resolve().relative_to(paths.ROOT).as_posix()


def clean(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (np.integer, np.floating)):
        value = value.item()
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    return value


def compute_midrank(values: np.ndarray) -> np.ndarray:
    order = np.argsort(values)
    sorted_values = values[order]
    ranks = np.zeros(len(values), dtype=float)
    start = 0
    while start < len(values):
        end = start
        while end < len(values) and sorted_values[end] == sorted_values[start]:
            end += 1
        ranks[start:end] = 0.5 * (start + end - 1) + 1
        start = end
    result = np.empty(len(values), dtype=float)
    result[order] = ranks
    return result


def delong_auc_ci(y_true: np.ndarray, scores: np.ndarray) -> tuple[float, float, float]:
    """Return one-vs-rest AUC and a naive image-level DeLong 95% CI."""
    y_true = np.asarray(y_true, dtype=int)
    scores = np.asarray(scores, dtype=float)
    positive = scores[y_true == 1]
    negative = scores[y_true == 0]
    m, n = len(positive), len(negative)
    if m == 0 or n == 0:
        return math.nan, math.nan, math.nan
    tx = compute_midrank(positive)
    ty = compute_midrank(negative)
    tz = compute_midrank(np.r_[positive, negative])
    auc = (tz[:m].sum() / m - (m + 1) / 2) / n
    v01 = (tz[:m] - tx) / n
    v10 = 1 - (tz[m:] - ty) / m
    sx = np.cov(v01, ddof=1) if m > 1 else 0.0
    sy = np.cov(v10, ddof=1) if n > 1 else 0.0
    standard_error = math.sqrt(max(float(sx / m + sy / n), 0.0))
    z = norm.ppf(0.975)
    return float(auc), max(0.0, auc - z * standard_error), min(1.0, auc + z * standard_error)


def wilson(k: int, n: int, z: float = 1.96) -> tuple[float, float, float]:
    if n == 0:
        return math.nan, math.nan, math.nan
    estimate = k / n
    denominator = 1 + z * z / n
    centre = (estimate + z * z / (2 * n)) / denominator
    half_width = z * math.sqrt(
        estimate * (1 - estimate) / n + z * z / (4 * n * n)
    ) / denominator
    return estimate, max(0.0, centre - half_width), min(1.0, centre + half_width)


def bootstrap_intervals(
    y_true: np.ndarray,
    y_pred: np.ndarray,
    probabilities: np.ndarray,
) -> dict[str, tuple[float, float]]:
    """Naive stratified image-level percentile intervals for global metrics."""
    rng = np.random.default_rng(BOOTSTRAP_SEED)
    class_indices = [np.flatnonzero(y_true == class_id) for class_id in range(3)]
    samples: dict[str, list[float]] = {
        "balanced_accuracy": [],
        "macro_f1_dice": [],
        "weighted_f1": [],
        "macro_auc_ovr": [],
        "weighted_auc_ovr": [],
        "multiclass_log_loss": [],
        "multiclass_brier": [],
    }
    for _ in range(BOOTSTRAP_REPLICATES):
        indices = np.concatenate(
            [rng.choice(index, size=len(index), replace=True) for index in class_indices]
        )
        yt = y_true[indices]
        yp = y_pred[indices]
        pr = probabilities[indices]
        samples["balanced_accuracy"].append(balanced_accuracy_score(yt, yp))
        samples["macro_f1_dice"].append(f1_score(yt, yp, average="macro"))
        samples["weighted_f1"].append(f1_score(yt, yp, average="weighted"))
        samples["macro_auc_ovr"].append(
            roc_auc_score(yt, pr, multi_class="ovr", average="macro")
        )
        samples["weighted_auc_ovr"].append(
            roc_auc_score(yt, pr, multi_class="ovr", average="weighted")
        )
        samples["multiclass_log_loss"].append(log_loss(yt, pr, labels=[0, 1, 2]))
        one_hot = np.eye(3, dtype=float)[yt]
        samples["multiclass_brier"].append(float(np.mean(np.sum((pr - one_hot) ** 2, axis=1))))
    return {
        name: (float(np.quantile(values, 0.025)), float(np.quantile(values, 0.975)))
        for name, values in samples.items()
    }


def calibration_table(
    y_true: np.ndarray,
    y_pred: np.ndarray,
    probabilities: np.ndarray,
) -> tuple[pd.DataFrame, dict[str, float]]:
    records: list[dict[str, Any]] = []
    summaries: dict[str, float] = {}
    definitions = [
        ("top_label", probabilities.max(axis=1), (y_true == y_pred).astype(float)),
        ("melanoma_ovr", probabilities[:, 1], (y_true == 1).astype(float)),
    ]
    for calibration_type, confidence, observed in definitions:
        bin_ids = np.minimum((confidence * 10).astype(int), 9)
        weighted_gap = 0.0
        maximum_gap = 0.0
        for bin_id in range(10):
            mask = bin_ids == bin_id
            count = int(mask.sum())
            mean_confidence = float(confidence[mask].mean()) if count else math.nan
            observed_rate = float(observed[mask].mean()) if count else math.nan
            gap = abs(mean_confidence - observed_rate) if count else math.nan
            if count:
                weighted_gap += count / len(y_true) * gap
                maximum_gap = max(maximum_gap, gap)
            records.append(
                {
                    "calibration_type": calibration_type,
                    "bin": bin_id + 1,
                    "lower_inclusive": bin_id / 10,
                    "upper_inclusive_for_last": (bin_id + 1) / 10,
                    "n": count,
                    "mean_predicted_probability": mean_confidence,
                    "observed_frequency": observed_rate,
                    "absolute_gap": gap,
                }
            )
        summaries[f"ece_10bin_{calibration_type}"] = weighted_gap
        summaries[f"mce_10bin_{calibration_type}"] = maximum_gap
    return pd.DataFrame.from_records(records), summaries


def metric_record(
    scope: str,
    metric: str,
    estimate: float | int,
    *,
    class_name: str = "all",
    ci_low: float = math.nan,
    ci_high: float = math.nan,
    ci_method: str = "none",
    denominator: int | float | None = None,
    note: str = "",
) -> dict[str, Any]:
    return {
        "scope": scope,
        "class": class_name,
        "metric": metric,
        "estimate": estimate,
        "ci95_low": ci_low,
        "ci95_high": ci_high,
        "ci_method": ci_method,
        "denominator_n": denominator,
        "note": note,
    }


def load_and_audit() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    required = [
        PREDICTIONS_INPUT,
        REGISTRY_INPUT,
        METADATA_INPUT,
        RAW_MANIFEST_INPUT,
        MODEL_MANIFEST_INPUT,
        PERFORMANCE_REGISTRY_INPUT,
        CURRENT_FEATURE_IMPORTANCE_INPUT,
        LEGACY_FEATURE_IMPORTANCE_INPUT,
        HANDCRAFT_ABLATION_INPUT,
    ]
    missing = [relative(path) for path in required if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing required input(s): " + ", ".join(missing))

    all_predictions = pd.read_csv(PREDICTIONS_INPUT)
    registry = pd.read_csv(REGISTRY_INPUT)
    metadata = pd.read_csv(METADATA_INPUT)
    manifest = pd.read_csv(RAW_MANIFEST_INPUT)
    cohort = all_predictions.loc[
        (all_predictions["condition"] == "1a")
        & (all_predictions["source"] == "ISIC2019")
        & (all_predictions["split_original"] == "test")
    ].copy()
    raw_cohort = manifest.loc[
        (manifest["source_reconstructed"] == "ISIC2019")
        & (manifest["split_original"] == "test")
    ].copy()

    if len(cohort) != 2_239:
        raise ValueError(f"Expected 2,239 ISIC2019/test predictions, found {len(cohort):,}")
    expected_counts = {"nevus": 894, "melanoma": 542, "atypical": 803}
    if cohort["class_name"].value_counts().to_dict() != expected_counts:
        raise ValueError("Unexpected class counts in ISIC2019/test predictions")
    if len(raw_cohort) != len(cohort):
        raise ValueError("Prediction and raw-manifest cohort sizes differ")
    if cohort["image_id"].duplicated().any() or raw_cohort["image_id"].duplicated().any():
        raise ValueError("Duplicate image_id within ISIC2019/test")
    if raw_cohort["sha256"].duplicated().any():
        raise ValueError("Duplicate raw SHA-256 within ISIC2019/test")
    if not np.allclose(cohort[PROBABILITY_COLUMNS].sum(axis=1), 1.0, atol=1e-6):
        raise ValueError("Probabilities do not sum to one")
    if not np.isfinite(cohort[PROBABILITY_COLUMNS].to_numpy(float)).all():
        raise ValueError("Non-finite probabilities found")

    train = manifest.loc[manifest["split_original"] == "train"]
    train_valid = manifest.loc[manifest["split_original"].isin(["train", "valid"])]
    overlap_checks = {
        "image_id_vs_train": set(raw_cohort["image_id"]) & set(train["image_id"]),
        "sha256_vs_train": set(raw_cohort["sha256"]) & set(train["sha256"]),
        "image_id_vs_train_valid": set(raw_cohort["image_id"]) & set(train_valid["image_id"]),
        "sha256_vs_train_valid": set(raw_cohort["sha256"]) & set(train_valid["sha256"]),
    }
    failures = {name: values for name, values in overlap_checks.items() if values}
    if failures:
        raise ValueError(
            "Held-out cohort overlaps fitting data: "
            + ", ".join(f"{name}={len(values)}" for name, values in failures.items())
        )

    registry_1a = registry.loc[registry["condition"] == "1a"]
    if len(registry_1a) != 1:
        raise ValueError("Expected exactly one 1a experiment-registry row")
    return cohort, raw_cohort, metadata, registry_1a


def build_prediction_export(cohort: pd.DataFrame, raw_cohort: pd.DataFrame) -> pd.DataFrame:
    raw_columns = [
        "image_id",
        "relative_path",
        "file",
        "bytes",
        "sha256",
        "class_harmonized",
    ]
    raw = raw_cohort[raw_columns].rename(
        columns={
            "relative_path": "raw_relative_path",
            "file": "raw_file",
            "bytes": "raw_bytes",
            "sha256": "raw_sha256",
            "class_harmonized": "raw_class_harmonized",
        }
    )
    result = cohort.merge(raw, on="image_id", how="left", validate="one_to_one")
    if result["raw_sha256"].isna().any():
        raise ValueError("Some prediction rows have no raw-manifest match")
    if not (result["class_name"] == result["raw_class_harmonized"]).all():
        raise ValueError("Prediction labels and raw-manifest labels differ")

    probabilities = result[PROBABILITY_COLUMNS].to_numpy(float)
    sorted_probabilities = np.sort(probabilities, axis=1)
    result["evaluation_cohort"] = "ISIC2019/test held-out"
    result["model_input_file"] = result["file"]
    result["true_class"] = result["class_name"]
    result["true_class_id"] = result["class_id"].astype(int)
    result["predicted_class_argmax"] = result["argmax_class_name"]
    result["predicted_class_id_argmax"] = result["argmax_class_id"].astype(int)
    result["correct_argmax"] = (
        result["true_class_id"] == result["predicted_class_id_argmax"]
    ).astype(int)
    result["max_probability"] = probabilities.max(axis=1)
    result["probability_margin_top2"] = sorted_probabilities[:, -1] - sorted_probabilities[:, -2]
    result["true_melanoma"] = (result["true_class"] == "melanoma").astype(int)
    result["error_transition"] = np.where(
        result["correct_argmax"].eq(1),
        "correct",
        result["true_class"] + " -> " + result["predicted_class_argmax"],
    )
    columns = [
        "evaluation_cohort",
        "condition",
        "row_index",
        "image_id",
        "raw_file",
        "raw_relative_path",
        "raw_bytes",
        "raw_sha256",
        "model_input_file",
        "true_class",
        "true_class_id",
        "p_nevus",
        "p_melanoma",
        "p_atypical",
        "probability_sum",
        "predicted_class_argmax",
        "predicted_class_id_argmax",
        "correct_argmax",
        "max_probability",
        "probability_margin_top2",
        "threshold_melanoma",
        "true_melanoma",
        "predicted_melanoma_at_threshold",
        "error_transition",
        "group_key",
        "fold",
    ]
    return result[columns].sort_values("row_index").reset_index(drop=True)


def build_metrics(predictions: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    y_true = predictions["true_class_id"].to_numpy(int)
    y_pred = predictions["predicted_class_id_argmax"].to_numpy(int)
    probabilities = predictions[PROBABILITY_COLUMNS].to_numpy(float)
    one_hot = np.eye(3, dtype=float)[y_true]
    intervals = bootstrap_intervals(y_true, y_pred, probabilities)
    calibration, calibration_summaries = calibration_table(y_true, y_pred, probabilities)

    records: list[dict[str, Any]] = []
    correct = int((y_true == y_pred).sum())
    accuracy, accuracy_low, accuracy_high = wilson(correct, len(y_true))
    records.extend(
        [
            metric_record("global_argmax", "n_images", len(y_true), denominator=len(y_true)),
            metric_record("global_argmax", "n_correct", correct, denominator=len(y_true)),
            metric_record(
                "global_argmax",
                "accuracy",
                accuracy,
                ci_low=accuracy_low,
                ci_high=accuracy_high,
                ci_method="Wilson 95%, image level",
                denominator=len(y_true),
            ),
        ]
    )
    global_values = {
        "balanced_accuracy": balanced_accuracy_score(y_true, y_pred),
        "macro_f1_dice": f1_score(y_true, y_pred, average="macro"),
        "weighted_f1": f1_score(y_true, y_pred, average="weighted"),
        "macro_auc_ovr": roc_auc_score(
            y_true, probabilities, multi_class="ovr", average="macro"
        ),
        "weighted_auc_ovr": roc_auc_score(
            y_true, probabilities, multi_class="ovr", average="weighted"
        ),
        "multiclass_log_loss": log_loss(y_true, probabilities, labels=[0, 1, 2]),
        "multiclass_brier": float(np.mean(np.sum((probabilities - one_hot) ** 2, axis=1))),
    }
    for metric, estimate in global_values.items():
        low, high = intervals[metric]
        records.append(
            metric_record(
                "global_argmax",
                metric,
                estimate,
                ci_low=low,
                ci_high=high,
                ci_method=f"stratified percentile bootstrap, {BOOTSTRAP_REPLICATES} replicates, seed {BOOTSTRAP_SEED}",
                denominator=len(y_true),
                note="Naive image-level interval; clustering by patient/lesion unavailable.",
            )
        )
    for metric, estimate in calibration_summaries.items():
        records.append(
            metric_record(
                "calibration",
                metric,
                estimate,
                denominator=len(y_true),
                note="10 equal-width probability bins.",
            )
        )

    confusion = confusion_matrix(y_true, y_pred, labels=[0, 1, 2])
    for class_id, class_name in enumerate(CLASS_ORDER):
        binary_true = (y_true == class_id).astype(int)
        binary_pred = (y_pred == class_id).astype(int)
        tn, fp, fn, tp = confusion_matrix(binary_true, binary_pred, labels=[0, 1]).ravel()
        auc, auc_low, auc_high = delong_auc_ci(binary_true, probabilities[:, class_id])
        sensitivity, sensitivity_low, sensitivity_high = wilson(int(tp), int(tp + fn))
        specificity, specificity_low, specificity_high = wilson(int(tn), int(tn + fp))
        ppv, ppv_low, ppv_high = wilson(int(tp), int(tp + fp))
        npv, npv_low, npv_high = wilson(int(tn), int(tn + fn))
        class_records = [
            ("support", int(tp + fn), math.nan, math.nan, "none", int(tp + fn)),
            ("tp", int(tp), math.nan, math.nan, "none", int(tp + fn)),
            ("fn", int(fn), math.nan, math.nan, "none", int(tp + fn)),
            ("tn", int(tn), math.nan, math.nan, "none", int(tn + fp)),
            ("fp", int(fp), math.nan, math.nan, "none", int(tn + fp)),
            ("sensitivity_recall", sensitivity, sensitivity_low, sensitivity_high, "Wilson 95%, image level", int(tp + fn)),
            ("specificity", specificity, specificity_low, specificity_high, "Wilson 95%, image level", int(tn + fp)),
            ("precision_ppv", ppv, ppv_low, ppv_high, "Wilson 95%, image level", int(tp + fp)),
            ("npv", npv, npv_low, npv_high, "Wilson 95%, image level", int(tn + fn)),
            ("f1_dice", f1_score(binary_true, binary_pred), math.nan, math.nan, "none", len(y_true)),
            ("auc_ovr", auc, auc_low, auc_high, "DeLong 95%, image level", len(y_true)),
            ("average_precision_ovr", average_precision_score(binary_true, probabilities[:, class_id]), math.nan, math.nan, "none", len(y_true)),
            ("brier_ovr", brier_score_loss(binary_true, probabilities[:, class_id]), math.nan, math.nan, "none", len(y_true)),
        ]
        for metric, estimate, low, high, method, denominator in class_records:
            records.append(
                metric_record(
                    "per_class_argmax_ovr",
                    metric,
                    estimate,
                    class_name=class_name,
                    ci_low=low,
                    ci_high=high,
                    ci_method=method,
                    denominator=denominator,
                    note="One-vs-rest for AUC/AP/Brier and diagnostic rates.",
                )
            )

    threshold = float(predictions["threshold_melanoma"].iloc[0])
    if not np.allclose(predictions["threshold_melanoma"], threshold):
        raise ValueError("Multiple melanoma thresholds found in held-out cohort")
    melanoma_true = predictions["true_melanoma"].to_numpy(int)
    melanoma_pred = predictions["predicted_melanoma_at_threshold"].to_numpy(int)
    tn, fp, fn, tp = confusion_matrix(melanoma_true, melanoma_pred, labels=[0, 1]).ravel()
    sensitivity, sensitivity_low, sensitivity_high = wilson(int(tp), int(tp + fn))
    specificity, specificity_low, specificity_high = wilson(int(tn), int(tn + fp))
    ppv, ppv_low, ppv_high = wilson(int(tp), int(tp + fp))
    npv, npv_low, npv_high = wilson(int(tn), int(tn + fn))
    threshold_note = (
        "Threshold selected on validation predictions after the heads had been fit on train+valid; "
        "the threshold operating point is descriptive and not independently calibrated."
    )
    threshold_records = [
        ("threshold", threshold, math.nan, math.nan, "none", len(y_true)),
        ("tp", int(tp), math.nan, math.nan, "none", int(tp + fn)),
        ("fn", int(fn), math.nan, math.nan, "none", int(tp + fn)),
        ("tn", int(tn), math.nan, math.nan, "none", int(tn + fp)),
        ("fp", int(fp), math.nan, math.nan, "none", int(tn + fp)),
        ("sensitivity", sensitivity, sensitivity_low, sensitivity_high, "Wilson 95%, image level", int(tp + fn)),
        ("specificity", specificity, specificity_low, specificity_high, "Wilson 95%, image level", int(tn + fp)),
        ("precision_ppv", ppv, ppv_low, ppv_high, "Wilson 95%, image level", int(tp + fp)),
        ("npv", npv, npv_low, npv_high, "Wilson 95%, image level", int(tn + fn)),
        ("accuracy", accuracy_score(melanoma_true, melanoma_pred), math.nan, math.nan, "none", len(y_true)),
        ("balanced_accuracy", balanced_accuracy_score(melanoma_true, melanoma_pred), math.nan, math.nan, "none", len(y_true)),
        ("f1_dice", f1_score(melanoma_true, melanoma_pred), math.nan, math.nan, "none", len(y_true)),
    ]
    for metric, estimate, low, high, method, denominator in threshold_records:
        records.append(
            metric_record(
                "melanoma_fixed_threshold",
                metric,
                estimate,
                class_name="melanoma",
                ci_low=low,
                ci_high=high,
                ci_method=method,
                denominator=denominator,
                note=threshold_note,
            )
        )

    confusion_records = []
    for true_id, true_name in enumerate(CLASS_ORDER):
        row_total = int(confusion[true_id].sum())
        for predicted_id, predicted_name in enumerate(CLASS_ORDER):
            count = int(confusion[true_id, predicted_id])
            confusion_records.append(
                {
                    "true_class": true_name,
                    "predicted_class": predicted_name,
                    "n": count,
                    "row_percentage": count / row_total,
                }
            )
    return pd.DataFrame.from_records(records), pd.DataFrame(confusion_records), calibration


def build_audit(
    predictions: pd.DataFrame,
    raw_cohort: pd.DataFrame,
    metadata: pd.DataFrame,
) -> pd.DataFrame:
    manifest = pd.read_csv(RAW_MANIFEST_INPUT)
    train = manifest.loc[manifest["split_original"] == "train"]
    train_valid = manifest.loc[manifest["split_original"].isin(["train", "valid"])]
    checks: list[dict[str, Any]] = []

    def add(check: str, observed: Any, expected: Any, status: str, interpretation: str) -> None:
        checks.append(
            {
                "check": check,
                "observed": observed,
                "expected": expected,
                "status": status,
                "interpretation": interpretation,
            }
        )

    add("cohort_source", "ISIC2019", "ISIC2019", "PASS", "Evaluation source is fixed.")
    add("cohort_split", "test", "test", "PASS", "Only the original held-out test split is evaluated.")
    add("n_predictions", len(predictions), 2239, "PASS", "One row per image_id.")
    add("n_unique_image_ids", predictions["image_id"].nunique(), len(predictions), "PASS", "No duplicate image identifier in evaluation.")
    add("n_unique_raw_sha256", predictions["raw_sha256"].nunique(), len(predictions), "PASS", "No byte-identical duplicate within evaluation.")
    for class_name, expected in (("nevus", 894), ("melanoma", 542), ("atypical", 803)):
        observed = int((predictions["true_class"] == class_name).sum())
        add(f"n_{class_name}", observed, expected, "PASS", "Frozen class count.")
    overlap_specs = [
        ("image_id_overlap_vs_train", set(raw_cohort["image_id"]) & set(train["image_id"])),
        ("sha256_overlap_vs_train", set(raw_cohort["sha256"]) & set(train["sha256"])),
        ("image_id_overlap_vs_train_valid", set(raw_cohort["image_id"]) & set(train_valid["image_id"])),
        ("sha256_overlap_vs_train_valid", set(raw_cohort["sha256"]) & set(train_valid["sha256"])),
    ]
    for name, overlap in overlap_specs:
        add(name, len(overlap), 0, "PASS" if not overlap else "FAIL", "Exact identifier/hash overlap count.")
    add(
        "probability_sum_max_abs_error",
        float(np.max(np.abs(predictions[PROBABILITY_COLUMNS].sum(axis=1) - 1.0))),
        "<= 1e-6",
        "PASS",
        "Scores are normalized multiclass probabilities.",
    )
    add(
        "patient_lesion_independence",
        "not verifiable",
        "patient/lesion identifiers",
        "CAUTION",
        "Source files expose image identifiers only; exact patient/lesion independence cannot be proved.",
    )
    add(
        "validation_excluded_from_primary_eval",
        int(
            len(
                metadata.loc[
                    (metadata["source"] == "ISIC2019")
                    & (metadata["split_original"] == "valid")
                ]
            )
        ),
        "excluded",
        "PASS",
        "2,426 validation images are not reported as held-out evaluation because validation influenced selection/calibration.",
    )
    return pd.DataFrame.from_records(checks)


def add_title(ws, title: str, subtitle: str, end_col: int) -> int:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=20, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 32
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_col)
    cell = ws.cell(2, 1, subtitle)
    cell.font = Font(size=10, italic=True, color="404040")
    cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34
    return 4


def add_section(ws, row: int, title: str, end_col: int) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_col)
    cell = ws.cell(row, 1, title)
    cell.font = Font(size=12, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=TEAL)
    ws.row_dimensions[row].height = 22
    return row + 1


def add_key_value(ws, row: int, key: str, value: Any, note: str = "") -> int:
    ws.cell(row, 1, key)
    ws.cell(row, 2, clean(value))
    ws.cell(row, 1).font = Font(bold=True, color=NAVY)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=LIGHT_TEAL)
    if note:
        ws.merge_cells(start_row=row, start_column=3, end_row=row, end_column=8)
        ws.cell(row, 3, note)
        ws.cell(row, 3).font = Font(italic=True, color="666666")
        ws.cell(row, 3).alignment = Alignment(wrap_text=True, vertical="top")
    return row + 1


def add_table(
    ws,
    headers: list[str],
    rows: Iterable[Iterable[Any]],
    *,
    start_row: int,
    name: str,
    style: str = "TableStyleMedium2",
) -> int:
    materialized = [[clean(value) for value in row] for row in rows]
    for column, header in enumerate(headers, start=1):
        cell = ws.cell(start_row, column, header)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=THIN_GREY)
    ws.row_dimensions[start_row].height = 30
    for row_offset, values in enumerate(materialized, start=1):
        for column, value in enumerate(values, start=1):
            cell = ws.cell(start_row + row_offset, column, value)
            cell.alignment = Alignment(wrap_text=False, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
            if isinstance(value, float):
                cell.number_format = "0.000000"
            elif isinstance(value, int):
                cell.number_format = "#,##0"
    end_row = start_row + len(materialized)
    if materialized:
        reference = f"A{start_row}:{get_column_letter(len(headers))}{end_row}"
        table = Table(displayName=name, ref=reference)
        table.tableStyleInfo = TableStyleInfo(
            name=style,
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(table)
    return end_row + 2


def fit_columns(ws, min_width: int = 10, max_width: int = 45) -> None:
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        width = min_width
        for cell in column_cells:
            if cell.value is not None:
                width = max(width, min(len(str(cell.value)) + 2, max_width))
        ws.column_dimensions[letter].width = width


def add_dataframe_sheet(
    workbook: Workbook,
    name: str,
    title: str,
    subtitle: str,
    data: pd.DataFrame,
    table_name: str,
    *,
    max_width: int = 45,
) -> None:
    ws = workbook.create_sheet(name)
    start = add_title(ws, title, subtitle, max(8, len(data.columns)))
    add_table(
        ws,
        list(data.columns),
        data.itertuples(index=False, name=None),
        start_row=start,
        name=table_name,
    )
    ws.freeze_panes = f"A{start + 1}"
    ws.auto_filter.ref = ws.tables[table_name].ref
    fit_columns(ws, max_width=max_width)


def build_workbook(
    predictions: pd.DataFrame,
    metrics: pd.DataFrame,
    confusion: pd.DataFrame,
    calibration: pd.DataFrame,
    audit: pd.DataFrame,
    metadata: pd.DataFrame,
    registry: pd.DataFrame,
) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    model_manifest = json.loads(MODEL_MANIFEST_INPUT.read_text(encoding="utf-8"))
    performance_registry = pd.read_csv(PERFORMANCE_REGISTRY_INPUT)
    current_importance = pd.read_csv(CURRENT_FEATURE_IMPORTANCE_INPUT)
    legacy_importance = pd.read_csv(LEGACY_FEATURE_IMPORTANCE_INPUT)
    handcraft_ablation = pd.read_csv(HANDCRAFT_ABLATION_INPUT)

    def select_metric(scope: str, metric: str, class_name: str = "all") -> pd.Series:
        selected = metrics.loc[
            (metrics["scope"] == scope)
            & (metrics["metric"] == metric)
            & (metrics["class"] == class_name)
        ]
        if len(selected) != 1:
            raise ValueError(f"Expected one metric row: {scope}/{class_name}/{metric}")
        return selected.iloc[0]

    # Sheet 1 — dense executive summary.
    ws = workbook.create_sheet("Executive Summary")
    row = add_title(
        ws,
        "SkinFusionNet — Complete Scientific Results",
        "Primary evaluation: 2,239 ISIC2019/test images absent from train and validation by exact image identifier and raw SHA-256. Historical results are included with explicit evidence status.",
        12,
    )
    row = add_section(ws, row, "Primary held-out evaluation", 12)
    for key, value, note in [
        ("Evaluation cohort", "ISIC2019 / original test split", "Nevus=894; melanoma=542; atypical=803."),
        ("Evaluated configuration", "1a rerun — V2S single-pass + LightGBM two-view ensemble", "Configuration-level result; not proof of byte identity with the final deployed LightGBM heads."),
        ("Exact overlap with fitting data", "0 image IDs; 0 raw SHA-256 hashes", "Checked against train and against train+validation."),
        ("Statistical unit", "image/file", "Patient and lesion identifiers are unavailable; intervals are naive image-level estimates."),
        ("Workbook scope", "results, architecture, model history, feature inspection, predictions, audit", "Raw images are not embedded."),
    ]:
        row = add_key_value(ws, row, key, value, note)
    row += 1
    row = add_section(ws, row, "Headline results", 12)
    headline_specs = [
        ("Accuracy", select_metric("global_argmax", "accuracy")),
        ("Balanced accuracy", select_metric("global_argmax", "balanced_accuracy")),
        ("Macro F1 / Dice", select_metric("global_argmax", "macro_f1_dice")),
        ("Macro AUC OVR", select_metric("global_argmax", "macro_auc_ovr")),
        ("Melanoma AUC OVR", select_metric("per_class_argmax_ovr", "auc_ovr", "melanoma")),
        ("Melanoma argmax sensitivity", select_metric("per_class_argmax_ovr", "sensitivity_recall", "melanoma")),
        ("Melanoma argmax specificity", select_metric("per_class_argmax_ovr", "specificity", "melanoma")),
    ]
    headline_rows = [
        [label, metric["estimate"], metric["ci95_low"], metric["ci95_high"], metric["ci_method"]]
        for label, metric in headline_specs
    ]
    row = add_table(
        ws,
        ["Metric", "Estimate", "95% CI low", "95% CI high", "Interval method"],
        headline_rows,
        start_row=row,
        name="HeadlineResults",
    )
    row = add_section(ws, row, "Model at a glance", 12)
    model_glance = [
        ("Bundle", model_manifest["bundle_id"], model_manifest["status"]),
        ("Task", "three-class dermoscopic image classification", ", ".join(model_manifest["classes"])),
        ("Views", "U-Net lesion-only + DullRazor dehaired", "one CNN pass per view"),
        ("Deep encoders", "2 × EfficientNetV2-S + CBAM", "1,280 latent features per view"),
        ("Handcrafted block", "18 ABCD-like descriptors", "shared with each fusion head"),
        ("Fusion heads", "2 × StandardScaler + LightGBM", "1,298 inputs and 3 probabilities per head"),
        ("Ensemble", model_manifest["ensemble"], "normalized multiclass probabilities"),
        ("Runtime", "five ONNX components, local inference", model_manifest["runtime_reference"]["app_version"]),
    ]
    row = add_table(
        ws,
        ["Item", "Value", "Detail"],
        model_glance,
        start_row=row,
        name="ModelAtGlance",
    )
    row = add_section(ws, row, "Interpretation boundaries", 12)
    boundary_rows = [
        ("Primary result", "Use ISIC2019/test metrics in this workbook for the requested held-out base."),
        ("Model history", "CV metrics are descriptive and optimistic because CNN representations were partly in-sample."),
        ("Threshold metrics", "The 0.5383 melanoma threshold is descriptive: validation was used after heads were fit on train+validation."),
        ("Feature rankings", "Current ONNX rankings are split frequencies, not gain, SHAP, causal importance, or proof of marginal value."),
        ("Legacy rankings", "The 512-d feature ranking belongs to ResNet18-era models and must not be attributed to the deployed V2S model."),
        ("Clinical use", "This research workbook is not clinical validation and not a standalone diagnostic device evaluation."),
    ]
    add_table(
        ws,
        ["Boundary", "Required interpretation"],
        boundary_rows,
        start_row=row,
        name="InterpretationBoundaries",
    )
    ws.freeze_panes = "A4"
    fit_columns(ws, max_width=86)

    # Sheet 2 — architecture, training recipe, and data composition.
    ws = workbook.create_sheet("Model & Data")
    row = add_title(
        ws,
        "Model Architecture and Data",
        "Executable component contracts come from the deployed model manifest. Training-recipe fields come from the current scripts; exact checkpoint hashes and the training commit were not retained.",
        12,
    )
    row = add_section(ws, row, "End-to-end inference architecture", 12)
    pipeline_rows = [
        (1, "Input standardization", "Dermoscopic image", "Square crop and resize", "512×512 working image"),
        (2, "Segmentation", "Working image resized to 384×384", "U-Net / EfficientNet-B0; sigmoid mask threshold 0.5", "Binary lesion mask"),
        (3, "View A", "Image + mask", "Black background outside lesion", "Lesion-only view"),
        (4, "View C", "Working image", "DullRazor black-hat morphology and inpainting", "Dehaired view"),
        (5, "Handcrafted block", "Mask + lesion pixels", "ABCD-like geometry, colour, border and fractal descriptors", "18 shared variables"),
        (6, "Deep encoding A/C", "Each view resized to 224×224 and ImageNet-normalized", "EfficientNetV2-S + CBAM + global average pooling", "1,280 features per view"),
        (7, "Fusion heads A/C", "18 handcrafted + 1,280 deep", "StandardScaler + LightGBM", "3 normalized probabilities per head"),
        (8, "Ensemble", "P_A and P_C", "Arithmetic mean", "P(nevus), P(melanoma), P(atypical)"),
        (9, "Decision readout", "Normalized probabilities", "Argmax or explicit class threshold", "Class prediction / flags"),
    ]
    row = add_table(
        ws,
        ["Step", "Stage", "Input", "Operation", "Output"],
        pipeline_rows,
        start_row=row,
        name="InferenceArchitecture",
    )
    row = add_section(ws, row, "Deployed ONNX components", 12)
    component_rows = []
    for component in model_manifest["components"]:
        component_rows.append(
            [
                component["name"],
                component["role"],
                component["architecture"],
                component["input"],
                component["output"],
                component["onnx_opset"],
                component.get("onnx_ml_opset"),
                component["sha256"],
            ]
        )
    row = add_table(
        ws,
        ["Component", "Role", "Architecture", "Input contract", "Output contract", "ONNX opset", "ONNX-ML opset", "SHA-256"],
        component_rows,
        start_row=row,
        name="DeployedComponents",
    )
    row = add_section(ws, row, "Training and model recipe recorded in code", 12)
    recipe_rows = [
        ("CNN initialization", "EfficientNetV2-S ImageNet-1K V1", "derm/pipeline/cnn_v2s.py"),
        ("Attention", "CBAM channel attention (reduction 16) then 7×7 spatial attention", "derm/legacy/pipeline/cnn.py"),
        ("CNN input / output", "224×224 RGB → 1,280-d penultimate embedding", "derm/pipeline/cnn_v2s.py"),
        ("CNN optimization", "12 epochs; AdamW; learning rate 1e-4; cosine annealing; cross-entropy", "derm/legacy/pipeline/cnn.py"),
        ("CNN batch size", "32 for EfficientNetV2-S", "derm/pipeline/cnn_v2s.py"),
        ("Augmentation", "horizontal/vertical flips; ±15° rotation; colour jitter 0.1", "derm/legacy/pipeline/cnn.py"),
        ("Checkpoint selection", "best validation accuracy", "derm/legacy/pipeline/cnn.py"),
        ("LightGBM", "300 estimators; learning rate 0.05; 63 leaves; random seed 42", "derm/fusion/train_lgbm_v2s.py"),
        ("Fusion scaling", "StandardScaler applied to each 1,298-d branch input", "model_manifest.json / ONNX Scaler nodes"),
        ("Known gap", "exact CNN checkpoint hashes, exact training commit, and exact historical dataset fingerprint were not recorded", "model_manifest.json"),
    ]
    row = add_table(
        ws,
        ["Parameter", "Recorded value", "Source"],
        recipe_rows,
        start_row=row,
        name="TrainingRecipe",
    )
    row = add_section(ws, row, "Merged dataset composition", 12)
    composition = (
        metadata.groupby(["source", "split_original", "class_name"])
        .size()
        .unstack(fill_value=0)
        .reset_index()
    )
    for class_name in CLASS_ORDER:
        if class_name not in composition.columns:
            composition[class_name] = 0
    composition["total"] = composition[list(CLASS_ORDER)].sum(axis=1)
    composition = composition[["source", "split_original", *CLASS_ORDER, "total"]]
    row = add_table(
        ws,
        list(composition.columns),
        composition.itertuples(index=False, name=None),
        start_row=row,
        name="DatasetComposition",
    )
    row = add_section(ws, row, "Handcrafted input dictionary", 12)
    handcraft_rows = [
        ("A", "A4", "Asymmetry", "Pearson correlation of horizontal and vertical lesion-mask projections"),
        ("B", "B", "Border compactness", "Perimeter² / (4π × area)"),
        ("D", "D", "Diameter", "Maximum Feret diameter on the convex hull, in pixels"),
        ("C", "Lab_a_kurt", "Colour", "Kurtosis of Lab a-channel lesion pixels"),
        ("C", "Lab_b_std", "Colour", "Standard deviation of Lab b-channel lesion pixels"),
        ("C", "Lab_a_std", "Colour", "Standard deviation of Lab a-channel lesion pixels"),
        ("C", "Lab_b_skew", "Colour", "Skewness of Lab b-channel lesion pixels"),
        ("B", "internal_R1..R16", "Inner-border morphology", "Five erosion-ring areas at radii 1, 2, 4, 8 and 16"),
        ("B", "external_R1..R16", "Outer-border morphology", "Five dilation-ring areas at radii 1, 2, 4, 8 and 16"),
        ("B", "fractal_D", "Border roughness", "Minkowski-Bouligand estimate from multi-radius boundary growth"),
    ]
    add_table(
        ws,
        ["ABCD block", "Feature(s)", "Concept", "Operational definition"],
        handcraft_rows,
        start_row=row,
        name="HandcraftedDictionary",
    )
    ws.freeze_panes = "A4"
    fit_columns(ws, max_width=82)

    # Sheet 3 — all primary metrics, confusion, and calibration in one place.
    ws = workbook.create_sheet("Performance")
    row = add_title(
        ws,
        "ISIC2019/test Performance",
        "All primary metrics are recalculated from the 2,239 observation-level predictions. Confidence intervals do not account for patient/lesion clustering because those identifiers are unavailable.",
        12,
    )
    row = add_section(ws, row, "Complete metric registry", 12)
    row = add_table(
        ws,
        list(metrics.columns),
        metrics.itertuples(index=False, name=None),
        start_row=row,
        name="CompleteMetricRegistry",
    )
    row = add_section(ws, row, "Conventional confusion matrix — argmax", 12)
    confusion_wide = (
        confusion.pivot(index="true_class", columns="predicted_class", values="n")
        .reindex(index=CLASS_ORDER, columns=CLASS_ORDER)
        .astype(int)
    )
    confusion_rows = []
    for true_class in CLASS_ORDER:
        values = confusion_wide.loc[true_class].tolist()
        confusion_rows.append([true_class, *values, sum(values)])
    confusion_rows.append(["predicted total", *confusion_wide.sum(axis=0).tolist(), int(confusion_wide.values.sum())])
    row = add_table(
        ws,
        ["True class", "Predicted nevus", "Predicted melanoma", "Predicted atypical", "True-class total"],
        confusion_rows,
        start_row=row,
        name="ConfusionMatrixWide",
    )
    row = add_section(ws, row, "Calibration bins", 12)
    row = add_table(
        ws,
        list(calibration.columns),
        calibration.itertuples(index=False, name=None),
        start_row=row,
        name="CalibrationBins",
    )
    row = add_section(ws, row, "Cross-protocol evidence registry", 12)
    add_table(
        ws,
        list(performance_registry.columns),
        performance_registry.itertuples(index=False, name=None),
        start_row=row,
        name="CrossProtocolEvidence",
    )
    ws.freeze_panes = "A5"
    fit_columns(ws, max_width=76)

    # Sheet 4 — model lineage plus current and legacy feature evidence.
    ws = workbook.create_sheet("History & Features")
    row = add_title(
        ws,
        "Model History and Feature Evidence",
        "Historical scores are retained for lineage comparison only. Protocol, representation leakage, deployment status, and feature-ranking semantics are shown explicitly.",
        13,
    )
    row = add_section(ws, row, "Model lineage", 13)
    history_rows = [
        ("early baseline; exact date not recorded", "ResNet18 + CBAM", "legacy", 512, 18, "single-pass two-view mean", "10-fold image CV", 0.9205, 0.9819, 0.9038, None, "descriptive; partially in-sample representations", "results/legacy/lgbm_entropy_log.txt"),
        ("historical intermediate; exact date not recorded", "V2S historical single-pass", "irreproducible", 1344, 18, "single-pass two-view mean", "10-fold image CV", 0.9470, 0.9909, 0.9370, None, "archive; feature dimension differs from current V2S", "results/legacy/lgbm_v2s_log.txt"),
        ("2026-06-04", "V2S + CBAM TTA×8", "former deployed", 1280, 18, "D4 TTA×8 then two-view mean", "10-fold image CV", 0.9528, 0.9919, 0.9437, 0.902512, "archived; internal CV optimistic; external file-level n=107", "results/legacy/lgbm_v2s_tta_log.txt"),
        ("2026-06-30", "V2S + CBAM single-pass", "deployed", 1280, 18, "single-pass two-view mean", "10-fold image CV", 0.9424, 0.9886, 0.9312, 0.931818, "current bundle; internal CV optimistic; external file-level n=107", "model_manifest.json / performance registry"),
        ("2026-07-21 workbook rerun", "V2S single-pass configuration 1a", "current held-out rerun", 1280, 18, "single-pass two-view mean", "ISIC2019/test only", float(select_metric("global_argmax", "accuracy")["estimate"]), float(select_metric("global_argmax", "macro_auc_ovr")["estimate"]), float(select_metric("global_argmax", "macro_f1_dice")["estimate"]), None, "primary workbook result; patient/lesion independence unresolved; configuration-level heads", paths.relative(paths.ISIC2019_METRICS)),
    ]
    row = add_table(
        ws,
        ["Recorded era", "Lineage", "Status", "Deep dimension", "Handcrafted dimension", "Inference", "Internal protocol", "Accuracy", "Macro AUC OVR", "Macro F1/Dice", "External melanoma AUC", "Evidence limitation", "Source"],
        history_rows,
        start_row=row,
        name="ModelLineage",
    )
    row = add_section(ws, row, "Current deployed heads — top 30 input split frequencies", 13)
    current_columns = [
        "rank_combined",
        "feature_index",
        "feature",
        "feature_family",
        "clinical_block",
        "splits_lesion_head_A",
        "splits_dehair_head_C",
        "splits_combined",
        "share_all_branch_splits",
        "importance_semantics",
    ]
    current_top = current_importance.head(30)
    row = add_table(
        ws,
        current_columns,
        current_top[current_columns].itertuples(index=False, name=None),
        start_row=row,
        name="CurrentTopFeatures",
    )
    row = add_section(ws, row, "All 18 handcrafted inputs in the current deployed heads", 13)
    current_handcraft = current_importance.loc[current_importance["feature_family"] == "handcrafted"]
    row = add_table(
        ws,
        current_columns,
        current_handcraft[current_columns].itertuples(index=False, name=None),
        start_row=row,
        name="CurrentHandcraftedFeatures",
    )
    row = add_section(ws, row, "Legacy 512-d model — recorded top 25 features", 13)
    legacy_top = legacy_importance.head(25).copy()
    legacy_top["feature_family"] = np.where(
        legacy_top["feature"].str.startswith("cnn_"), "CNN_embedding", "handcrafted"
    )
    legacy_top["status"] = "legacy 512-d ranking; not the deployed model"
    row = add_table(
        ws,
        list(legacy_top.columns),
        legacy_top.itertuples(index=False, name=None),
        start_row=row,
        name="LegacyTopFeatures",
    )
    row = add_section(ws, row, "Handcrafted-only factorial results — internal image-level CV", 13)
    ablation_columns = [
        "condition",
        "n_features_hc",
        "n_eval",
        "auc_melanoma_ovr",
        "auc_melanoma_ci95_low",
        "auc_melanoma_ci95_high",
        "accuracy",
        "macro_dice",
        "sensitivity_melanoma",
        "specificity_melanoma",
        "status",
    ]
    ablation_view = handcraft_ablation.sort_values("auc_melanoma_ovr", ascending=False)
    row = add_table(
        ws,
        ablation_columns,
        ablation_view[ablation_columns].itertuples(index=False, name=None),
        start_row=row,
        name="HandcraftFactorial",
    )
    row = add_section(ws, row, "Feature conclusions that the evidence supports", 13)
    feature_conclusions = [
        ("Current split frequency", "The four Lab colour statistics occupy ranks 1–4 and fractal_D rank 5 across the deployed ONNX heads."),
        ("Relative split volume", f"Handcrafted inputs account for {current_handcraft['splits_combined'].sum():,} of {current_importance['splits_combined'].sum():,} branch splits; split count does not measure unique marginal value."),
        ("Handcrafted-only discrimination", "All 18 handcrafted variables reach melanoma AUC 0.8038; colour alone reaches 0.7286 in the recorded image-level CV."),
        ("Deep comparison", "Handcrafted-only performance remains below the held-out melanoma AUC of the V2S configuration."),
        ("Unanswered causal question", "The completed factorial contains handcrafted-only conditions; clean Deep versus Deep+block marginal effects remain incomplete."),
        ("Legacy caution", "The old 512-d ranking also emphasized Lab colour, but it cannot be used as current-model importance."),
    ]
    add_table(
        ws,
        ["Evidence", "Supported statement"],
        feature_conclusions,
        start_row=row,
        name="FeatureConclusions",
    )
    ws.freeze_panes = "A5"
    fit_columns(ws, max_width=80)

    # Sheet 5 — observation-level handoff. Errors are filterable in the same table.
    ws = workbook.create_sheet("Predictions")
    row = add_title(
        ws,
        "Raw ISIC2019/test Predictions",
        "One row per held-out image. Filter correct_argmax=0 or error_transition<>correct to review all 495 errors without duplicating them on a separate sheet.",
        len(predictions.columns),
    )
    row = add_section(ws, row, "Observation-level results", len(predictions.columns))
    add_table(
        ws,
        list(predictions.columns),
        predictions.itertuples(index=False, name=None),
        start_row=row,
        name="PredictionsISIC2019",
    )
    prediction_columns = {
        name: index + 1 for index, name in enumerate(predictions.columns)
    }
    correct_letter = get_column_letter(prediction_columns["correct_argmax"])
    confidence_letter = get_column_letter(prediction_columns["max_probability"])
    prediction_last_row = row + len(predictions)
    ws.conditional_formatting.add(
        f"{correct_letter}{row + 1}:{correct_letter}{prediction_last_row}",
        CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", fgColor=RED)),
    )
    ws.conditional_formatting.add(
        f"{correct_letter}{row + 1}:{correct_letter}{prediction_last_row}",
        CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", fgColor=GREEN)),
    )
    ws.conditional_formatting.add(
        f"{confidence_letter}{row + 1}:{confidence_letter}{prediction_last_row}",
        ColorScaleRule(
            start_type="num",
            start_value=0,
            start_color=RED,
            mid_type="num",
            mid_value=0.5,
            mid_color=YELLOW,
            end_type="num",
            end_value=1,
            end_color=GREEN,
        ),
    )
    ws.freeze_panes = f"A{row + 1}"
    fit_columns(ws, max_width=50)

    # Sheet 6 — integrity checks, hashes, limitations, and field definitions.
    ws = workbook.create_sheet("Audit & Provenance")
    row = add_title(
        ws,
        "Audit, Provenance, and Data Dictionary",
        "This sheet contains the checks needed to interpret and reproduce the workbook. A PASS confirms the implemented exact check, not patient-level independence or clinical validity.",
        12,
    )
    row = add_section(ws, row, "Cohort and overlap audit", 12)
    row = add_table(
        ws,
        list(audit.columns),
        audit.itertuples(index=False, name=None),
        start_row=row,
        name="CohortAudit",
    )
    provenance_records = []
    for source_path, role in [
        (PREDICTIONS_INPUT, "Source observation-level predictions for configuration 1a"),
        (REGISTRY_INPUT, "Experimental definitions and caveats"),
        (METADATA_INPUT, "Reconstructed image source, split, and class"),
        (RAW_MANIFEST_INPUT, "Current raw-image paths, sizes, and SHA-256 hashes"),
        (MODEL_MANIFEST_INPUT, "Deployed bundle architecture, contracts, hashes, and provenance"),
        (PERFORMANCE_REGISTRY_INPUT, "Cross-protocol evidence registry"),
        (CURRENT_FEATURE_IMPORTANCE_INPUT, "Current ONNX head split-frequency inspection"),
        (LEGACY_FEATURE_IMPORTANCE_INPUT, "Legacy 512-d feature ranking"),
        (HANDCRAFT_ABLATION_INPUT, "Handcrafted-only factorial results"),
        (paths.FEATURES_CSV, "18 aligned handcrafted variables"),
        (paths.CNN_V2S_FEATS_CSV, "Single-pass lesion-view CNN representations"),
        (paths.CNN_V2S_DEHAIR_CSV, "Single-pass dehaired-view CNN representations"),
        (PREDICTIONS_OUTPUT, "Filtered ISIC2019/test predictions"),
        (METRICS_OUTPUT, "Recalculated long-format metric export"),
        (AUDIT_OUTPUT, "Integrity and overlap check export"),
    ]:
        provenance_records.append(
            {
                "path": relative(source_path),
                "exists": source_path.exists(),
                "bytes": source_path.stat().st_size if source_path.exists() else math.nan,
                "sha256": file_sha256(source_path) if source_path.exists() else "",
                "role": role,
            }
        )
    provenance = pd.DataFrame(provenance_records)
    row = add_section(ws, row, "File provenance", 12)
    row = add_table(
        ws,
        list(provenance.columns),
        provenance.itertuples(index=False, name=None),
        start_row=row,
        name="FileProvenance",
    )
    row = add_section(ws, row, "Known limitations", 12)
    limitations = [
        ("Patient/lesion independence", "Not verifiable because patient_id and lesion_id are unavailable in the merged internal dataset."),
        ("Historical training identity", "The exact training commit, original dataset fingerprint, and source CNN checkpoint hashes were not recorded."),
        ("Primary evaluated object", "The ISIC2019 result evaluates a rerun configuration using current matrices, not byte-identical final ONNX head outputs."),
        ("Threshold calibration", "The melanoma threshold was selected on validation after the heads were fit on train+validation."),
        ("Historical CV", "CNN representations were partially in-sample and the historical scaler was fit before folds; metrics are optimistic."),
        ("External evidence", "Only 107 folder-labelled files, including 19 melanoma-labelled files; near-duplicates and reference-standard provenance are unresolved."),
        ("Atypical class", "A heterogeneous harmonized category, not one clinical diagnosis."),
        ("Feature importance", "Split frequency is biased toward frequently splittable/correlated variables and does not establish marginal benefit."),
    ]
    row = add_table(
        ws,
        ["Limitation", "Implication"],
        limitations,
        start_row=row,
        name="KnownLimitations",
    )
    row = add_section(ws, row, "Data dictionary", 12)
    dictionary = [
        ("image_id", "Canonical image identifier; it does not guarantee patient or lesion identity."),
        ("raw_sha256", "SHA-256 hash of the raw image file in the merged dataset."),
        ("row_index", "Canonical index linking metadata and feature matrices."),
        ("true_class", "Harmonized reference class: nevus, melanoma, or atypical."),
        ("p_nevus / p_melanoma / p_atypical", "Normalized multiclass probabilities; sum approximately equals 1."),
        ("predicted_class_argmax", "Class associated with the largest probability."),
        ("correct_argmax", "1 when argmax equals the reference class; otherwise 0."),
        ("max_probability", "Largest of the three class probabilities."),
        ("probability_margin_top2", "Difference between the two largest probabilities."),
        ("threshold_melanoma", "Youden threshold derived from validation; calibration is not independent."),
        ("predicted_melanoma_at_threshold", "1 when p_melanoma is greater than or equal to threshold_melanoma."),
        ("AUC OVR", "ROC area under the curve for one class versus all other classes."),
        ("F1 / Dice", "2×TP/(2×TP+FP+FN); F1 and Dice are identical here."),
        ("multiclass_brier", "Mean summed squared probability error across the three classes."),
        ("ECE 10 bins", "Weighted mean absolute calibration gap over ten equal-width probability bins."),
        ("split frequency", "Number of LightGBM branch nodes using an input; not a causal effect or gain importance."),
        ("evidence grade B", "Recorded held-out result with provenance or reproduction gaps."),
        ("evidence grade C", "Exploratory file-level external or small paired evidence."),
        ("evidence grade D", "Historical, archived, or optimistic descriptive evidence."),
    ]
    add_table(
        ws,
        ["Field or metric", "Definition"],
        dictionary,
        start_row=row,
        name="DataDictionary",
    )
    ws.freeze_panes = "A5"
    fit_columns(ws, max_width=86)

    for sheet in workbook.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.auto_filter.ref = None
        for row_cells in sheet.iter_rows():
            for cell in row_cells:
                status_colour = {
                    "PASS": GREEN,
                    "CAUTION": YELLOW,
                    "FAIL": RED,
                    "deployed": GREEN,
                    "former deployed": YELLOW,
                    "legacy": "E7E6E6",
                    "irreproducible": ORANGE,
                }.get(str(cell.value))
                if status_colour:
                    cell.fill = PatternFill("solid", fgColor=status_colour)
                    cell.font = Font(bold=True, color="1F1F1F")
    WORKBOOK_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(WORKBOOK_OUTPUT)


def validate_workbook(predictions: pd.DataFrame, metrics: pd.DataFrame) -> None:
    workbook = load_workbook(WORKBOOK_OUTPUT, read_only=False, data_only=False)
    expected_sheets = [
        "Executive Summary",
        "Model & Data",
        "Performance",
        "History & Features",
        "Predictions",
        "Audit & Provenance",
    ]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(f"Unexpected workbook sheets: {workbook.sheetnames}")
    prediction_sheet = workbook["Predictions"]
    if prediction_sheet.max_row != len(predictions) + 5:
        raise ValueError("Workbook prediction row count is inconsistent")
    if "PredictionsISIC2019" not in prediction_sheet.tables:
        raise ValueError("Workbook prediction table is missing")
    if "CompleteMetricRegistry" not in workbook["Performance"].tables:
        raise ValueError("Workbook metric registry table is missing")
    if "ModelLineage" not in workbook["History & Features"].tables:
        raise ValueError("Workbook model-history table is missing")
    if "CurrentTopFeatures" not in workbook["History & Features"].tables:
        raise ValueError("Workbook current-feature table is missing")
    if len(metrics) < 50:
        raise ValueError("Metric export is unexpectedly small")
    workbook.close()


def main() -> None:
    cohort, raw_cohort, metadata, registry = load_and_audit()
    predictions = build_prediction_export(cohort, raw_cohort)
    metrics, confusion, calibration = build_metrics(predictions)
    audit = build_audit(predictions, raw_cohort, metadata)

    PREDICTIONS_OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    predictions.to_csv(PREDICTIONS_OUTPUT, index=False, float_format="%.10g")
    metrics.to_csv(METRICS_OUTPUT, index=False, float_format="%.10g")
    audit.to_csv(AUDIT_OUTPUT, index=False, float_format="%.10g")
    build_workbook(predictions, metrics, confusion, calibration, audit, metadata, registry)
    validate_workbook(predictions, metrics)

    accuracy = metrics.loc[
        (metrics["scope"] == "global_argmax") & (metrics["metric"] == "accuracy"),
        "estimate",
    ].iloc[0]
    macro_auc = metrics.loc[
        (metrics["scope"] == "global_argmax") & (metrics["metric"] == "macro_auc_ovr"),
        "estimate",
    ].iloc[0]
    melanoma_auc = metrics.loc[
        (metrics["scope"] == "per_class_argmax_ovr")
        & (metrics["class"] == "melanoma")
        & (metrics["metric"] == "auc_ovr"),
        "estimate",
    ].iloc[0]
    print(f"Wrote {relative(WORKBOOK_OUTPUT)} ({WORKBOOK_OUTPUT.stat().st_size:,} bytes)")
    print(f"Cohort: n={len(predictions):,}; errors={(predictions['correct_argmax'] == 0).sum():,}")
    print(f"Accuracy={accuracy:.6f}; macro AUC OVR={macro_auc:.6f}; melanoma AUC={melanoma_auc:.6f}")
    print(f"Workbook SHA-256: {file_sha256(WORKBOOK_OUTPUT)}")


if __name__ == "__main__":
    main()
